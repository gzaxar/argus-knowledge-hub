"""A file as Markdown, to be a document's body as it is ("Keep it as it is"): the structure the file has, kept —
a Word file's headings, lists, bold and tables; a spreadsheet's sheets as tables; an e-mail's subject, date and
message; a web page's headings, lists, tables and links — and, for text without structure (a PDF's, a plain text
file's), paragraphs put back together where lines were only wrapped, bullets and numbered lines as lists, and the
lines that look like headings as headings.

What the assistant reads to suggest fields stays files.extract (plain text with its origin, [page 2]).
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from email import policy
from email.parser import BytesParser
from xml.etree import ElementTree as ET

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
BULLET = re.compile(r"^\s*[•▪◦·●○■□➢►\-\*–—]\s+(.*)$")
NUMBERED = re.compile(r"^\s*(\d{1,3})[.)]\s+(.*)$")
SECTION = re.compile(r"^(\d{1,2}(?:\.\d{1,2}){0,3})\.?\s+(\S.{0,90})$")


class Unreadable(ValueError):
    pass


def _cell(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").replace("|", "\\|").strip()


def table(rows: list[list[str]]) -> str:
    """Rows as a Markdown table, the first as its header."""
    rows = [[_cell(c) for c in r] for r in rows if any((c or "").strip() for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    if width == 1:
        return "\n".join(r[0] for r in rows)
    return "\n".join(["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width,
                      *("| " + " | ".join(r) + " |" for r in rows[1:])])


# --------------------------------------------------------------------------- text without structure

def _heading(line: str) -> str | None:
    """'## …' for a line that looks like a heading: numbered (2.1 Scope), or short and in capitals."""
    if len(line) > 90 or line.endswith((".", ",", ";", ":")) and not SECTION.match(line):
        return None
    m = SECTION.match(line)
    if m and not line.endswith((".", ",", ";")) and re.match(r"[A-ZÀ-Ý]", m.group(2)):
        return "#" * min(2 + m.group(1).count("."), 4) + " " + line
    letters = re.sub(r"[^A-Za-zÀ-ÿ]", "", line)
    if len(letters) >= 4 and letters.isupper() and len(line.split()) <= 10:
        return "## " + line
    return None


def text_to_markdown(text: str, wrapped: bool = False) -> str:
    """Plain text as Markdown. `wrapped` (a PDF's text): a line that does not end a sentence and is followed by
    one that continues it is the same paragraph, and a word cut by a hyphen at the line's end is joined."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\f", "\n\n")
    out: list[str] = []
    for para in re.split(r"\n\s*\n", text):
        lines = [line.rstrip() for line in para.split("\n") if line.strip()]
        if not lines:
            continue
        blocks: list[str] = []
        current = ""
        for line in lines:
            stripped = line.strip()
            b, n = BULLET.match(stripped), NUMBERED.match(stripped)
            h = _heading(stripped) if (len(lines) == 1 or not current) else None
            if b or n or h:
                if current:
                    blocks.append(current)
                current = ""
                blocks.append(f"- {b.group(1)}" if b else f"{n.group(1)}. {n.group(2)}" if n else h)
                continue
            if current and wrapped:
                if current.endswith("-") and re.match(r"[a-zà-ÿ]", stripped):
                    current = current[:-1] + stripped                   # a word cut at the end of the line
                    continue
                if not re.search(r"[.!?:;]$", current) and re.match(r"[a-zà-ÿ0-9(]", stripped):
                    current += " " + stripped
                    continue
            if current and blocks and blocks[-1].startswith(("- ", "1. ")) and wrapped \
                    and re.match(r"[a-zà-ÿ]", stripped):
                blocks[-1] += " " + stripped                            # a list item wrapped on to the next line
                continue
            if current:
                blocks.append(current)
            current = stripped
        if current:
            blocks.append(current)
        # Consecutive plain lines of one paragraph keep their line breaks (an address, a signature).
        joined: list[str] = []
        for blk in blocks:
            if joined and not blk.startswith(("#", "- ")) and not NUMBERED.match(blk) and \
                    not joined[-1].startswith(("#", "- ")) and not NUMBERED.match(joined[-1]):
                joined[-1] += "  \n" + blk
            else:
                joined.append(blk)
        # Lists and headings stand apart from the text around them.
        para_md = ""
        for blk in joined:
            if not para_md:
                para_md = blk
            elif blk.startswith(("- ",)) or NUMBERED.match(blk):
                prev = para_md.rsplit("\n", 1)[-1]
                sep = "\n" if prev.startswith("- ") or NUMBERED.match(prev) else "\n\n"
                para_md += sep + blk
            else:
                para_md += "\n\n" + blk
        out.append(para_md)
    return re.sub(r"\n{3,}", "\n\n", "\n\n".join(out)).strip()


# --------------------------------------------------------------------------- Word

def _runs(p) -> str:
    parts = []
    for r in p.iter(f"{W}r"):
        text = "".join((t.text or "") for t in r.iter(f"{W}t"))
        if r.find(f"{W}tab") is not None:
            text = " " + text
        if r.find(f"{W}br") is not None:
            text += "  \n"
        if not text:
            continue
        rpr = r.find(f"{W}rPr")

        def on(tag):
            el = rpr.find(f"{W}{tag}") if rpr is not None else None
            return el is not None and el.get(f"{W}val") not in ("0", "false")
        stripped = text.strip()
        if stripped and on("b"):
            text = text.replace(stripped, f"**{stripped}**")
        elif stripped and on("i"):
            text = text.replace(stripped, f"*{stripped}*")
        parts.append(text)
    out = "".join(parts)
    return re.sub(r"\*\*\s*\*\*", "", re.sub(r"\*\*(\s*)\*\*", r"\1", out)).strip()


def _paragraph(p) -> str:
    ppr = p.find(f"{W}pPr")
    style = ""
    level = None
    if ppr is not None:
        st = ppr.find(f"{W}pStyle")
        style = (st.get(f"{W}val") if st is not None else "") or ""
        num = ppr.find(f"{W}numPr")
        if num is not None:
            lvl = num.find(f"{W}ilvl")
            level = int(lvl.get(f"{W}val", "0")) if lvl is not None else 0
    text = _runs(p)
    if not text:
        return ""
    m = re.match(r"(?i)(?:heading|titolo|berschrift|titre)\s*(\d)", style)
    if m:
        return "#" * min(int(m.group(1)), 6) + " " + text.replace("**", "")
    if style.lower() in ("title", "titolo"):
        return "# " + text.replace("**", "")
    if level is not None or "list" in style.lower() or "elenco" in style.lower():
        return "  " * (level or 0) + "- " + text
    return text


def docx_markdown(content: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            root = ET.fromstring(z.read("word/document.xml"))
    except Exception as exc:
        raise Unreadable(f"The Word file could not be opened: {exc}")
    body = root.find(f"{W}body")
    if body is None:
        raise Unreadable("The Word file has no text.")
    blocks: list[str] = []
    for el in body:
        if el.tag == f"{W}p":
            md = _paragraph(el)
            if md:
                blocks.append(md)
        elif el.tag == f"{W}tbl":
            rows = [[" ".join(_runs(p) for p in tc.iter(f"{W}p")) for tc in tr.findall(f"{W}tc")]
                    for tr in el.findall(f"{W}tr")]
            md = table(rows)
            if md:
                blocks.append(md)
    out = ""
    for blk in blocks:                            # list items stay together; everything else is its own block
        if out and blk.lstrip().startswith("- ") and out.rsplit("\n", 1)[-1].lstrip().startswith("- "):
            out += "\n" + blk
        else:
            out += ("\n\n" if out else "") + blk
    if not out.strip():
        raise Unreadable("The Word file has no text.")
    return out


# --------------------------------------------------------------------------- spreadsheets, CSV, e-mail

def xlsx_markdown(content: bytes, max_cells: int = 20000) -> str:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)   # cached values, never formulas
    except Exception as exc:
        raise Unreadable(f"The spreadsheet could not be opened: {exc}")
    parts, cells = [], 0
    visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
    for ws in visible:
        rows = []
        for row in ws.iter_rows(values_only=True):
            values = ["" if v is None else str(v) for v in row]
            while values and not values[-1].strip():
                values.pop()
            if any(v.strip() for v in values):
                rows.append(values)
                cells += len(values)
            if cells >= max_cells:
                break
        md = table(rows)
        if md:
            parts.append((f"## {ws.title}\n\n" if len(visible) > 1 else "") + md)
        if cells >= max_cells:
            break
    if not parts:
        raise Unreadable("The spreadsheet has no values.")
    return "\n\n".join(parts)


def csv_markdown(text: str) -> str:
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return table(list(csv.reader(io.StringIO(text), dialect)))


def email_markdown(content: bytes) -> str:
    """Its subject, its date and its message; not who sent it (as for the assistant: no address is passed on)."""
    msg = BytesParser(policy=policy.default).parsebytes(content)
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body is not None else ""
    if body is not None and body.get_content_type() == "text/html":
        md = html_markdown(text.encode("utf-8"))
    else:
        md = text_to_markdown(text or "")
    head = []
    if msg.get("subject"):
        head.append(f"**Subject:** {msg['subject']}")
    if msg.get("date"):
        head.append(f"**Date:** {msg['date']}")
    return ("  \n".join(head) + "\n\n" + md).strip()


def html_markdown(content: bytes) -> str:
    """One web page on its own: its pictures (not given with it) become their descriptions."""
    from app.services import web_import
    page = web_import.convert("page.html", {"page.html": content})
    md = re.sub(r"!\[([^\]]*)\]\(file:[^)]+\)", r"\1", page.markdown)
    return re.sub(r"\[([^\]]*)\]\(page:[^)]+\)", r"\1", md)
